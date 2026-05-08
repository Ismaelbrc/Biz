import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SOURCE = '/home/user/Biz/winds-of-wycaro.md'
OUTPUT = '/home/user/Biz/winds-of-wycaro-livro1.docx'

doc = Document()

section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.left_margin = Cm(3)
section.right_margin = Cm(3)
section.top_margin = Cm(3)
section.bottom_margin = Cm(2.5)

style_normal = doc.styles['Normal']
style_normal.font.name = 'Georgia'
style_normal.font.size = Pt(12)
style_normal.paragraph_format.line_spacing = Pt(20)
style_normal.paragraph_format.space_after = Pt(0)

# --- helpers ---

def add_inline(paragraph, text):
    parts = re.split(r'(\*[^*]+\*)', text)
    for part in parts:
        if part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        run.font.name = 'Georgia'
        run.font.size = Pt(12)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Georgia'
    run.font.size = Pt(26)
    run.bold = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Georgia'
    run.font.size = Pt(13)
    run.italic = True
    p.paragraph_format.space_after = Pt(40)

def add_part_title(text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Georgia'
    run.font.size = Pt(16)
    run.bold = True
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(60)

def add_chapter_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Georgia'
    run.font.size = Pt(16)
    run.bold = True
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(24)

def add_separator():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('* * *')
    run.font.name = 'Georgia'
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)

def add_pov_marker(name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.italic = True
    run.font.name = 'Georgia'
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)

def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(20)
    add_inline(p, text)

def add_fim(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.bold = True
    run.font.name = 'Georgia'
    run.font.size = Pt(14)
    p.paragraph_format.space_before = Pt(48)

# --- patterns ---

RE_TITLE     = re.compile(r'^#\s+WINDS OF WYCARO')
RE_SUBTITLE  = re.compile(r'^###\s+(.+)')
RE_PART      = re.compile(r'^#{1,2}\s+(PARTE\s+.+)', re.IGNORECASE)
RE_CHAPTER   = re.compile(r'^##\s+(Capítulo\s+\d+.+)')
RE_POV       = re.compile(r'^\*([A-Za-záéíóúàãõâêôçÁÉÍÓÚÀÃÕÂÊÔÇ]+)\*$')
RE_FIM       = re.compile(r'^\*FIM DO LIVRO')
RE_SEP       = re.compile(r'^---$')

# --- parse ---

with open(SOURCE, encoding='utf-8') as f:
    lines = [l.rstrip('\n') for l in f.readlines()]

chapter_started = False

for line in lines:
    stripped = line.strip()

    if not stripped:
        continue

    if RE_TITLE.match(stripped):
        add_title('WINDS OF WYCARO')
        continue

    m = RE_SUBTITLE.match(stripped)
    if m:
        add_subtitle(m.group(1))
        continue

    m = RE_PART.match(stripped)
    if m:
        add_part_title(m.group(1))
        continue

    m = RE_CHAPTER.match(stripped)
    if m:
        chapter_started = True
        add_chapter_title(m.group(1))
        continue

    if RE_SEP.match(stripped):
        if chapter_started:
            add_separator()
        continue

    if RE_FIM.match(stripped):
        text = stripped.strip('*').strip()
        add_fim(text)
        continue

    m = RE_POV.match(stripped)
    if m:
        add_pov_marker(m.group(1))
        continue

    add_body(stripped)

doc.save(OUTPUT)
print(f'Salvo: {OUTPUT}')
